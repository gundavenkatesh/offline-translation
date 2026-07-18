"""Document ingestion: extract -> chunk -> embed -> store in Chroma."""
import os
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_ollama import OllamaEmbeddings


CHROMA_DIR = os.environ.get("CHROMA_DIR", "data/chroma")
EMBED_MODEL = os.environ.get("EMBED_MODEL", "bge-m3")
# served by Ollama: `ollama pull bge-m3`
# alternatives: "nomic-embed-text" (English-leaning), "embeddinggemma"

_embeddings = None


def get_embeddings():
    global _embeddings
    if _embeddings is None:
        _embeddings = OllamaEmbeddings(model=EMBED_MODEL)
    return _embeddings


def get_vectorstore() -> Chroma:
    return Chroma(
        collection_name="documents",
        persist_directory=CHROMA_DIR,
        embedding_function=get_embeddings(),
    )


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        with fitz.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    if ext == ".docx":
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".md"):
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {ext}")


def ingest_file(path: str) -> int:
    """Extract, chunk, and store a file. Returns number of chunks added."""
    text = extract_text(path)
    if not text.strip():
        return 0

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_text(text)
    metadatas = [{"source": Path(path).name, "chunk": i} for i in range(len(chunks))]

    vs = get_vectorstore()
    vs.add_texts(chunks, metadatas=metadatas)
    return len(chunks)


def list_sources() -> list[str]:
    vs = get_vectorstore()
    data = vs.get(include=["metadatas"])
    return sorted({m["source"] for m in data["metadatas"] if m and "source" in m})
